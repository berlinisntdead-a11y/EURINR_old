"""
Collect EUR -> INR quotes from each provider and append them to the rates workbook
(data/EURINR_data.xlsx, sheet "Rates") and to a CSV history (data/eurinr_quotes.csv).

    python scraper/scrape_rates.py                      # all providers, default amounts
    python scraper/scrape_rates.py --providers wise,xe  # a subset
    python scraper/scrape_rates.py --dry-run --debug    # print raw responses, write nothing

Exits 1 when any provider failed, so a scheduled run shows up red and GitHub sends an email.
Rows from the providers that did succeed are still written.
"""
from __future__ import annotations

import argparse
import copy
import csv
import os
import sys
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

import openpyxl

import providers as P

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = ROOT / "data" / "EURINR_data.xlsx"
DEFAULT_CSV = ROOT / "data" / "eurinr_quotes.csv"
AMOUNTS = [100, 500, 1000, 2000, 5000]
MAX_DEVIATION = 0.10  # drop a quote whose rate is >10% from mid: almost certainly a parsing error
CET = ZoneInfo("Europe/Paris")

CSV_FIELDS = ["provider", "capture_date", "capture_time_cet", "payment_method", "amount_eur",
              "mid_rate", "recipient_inr", "fee_eur", "rate", "promo_rate", "fx_margin", "notes"]


def collect(names, amounts):
    http = P.Http()
    rows, failures = [], []
    for name in names:
        try:
            mid, mid_source = P.mid_market(http)
            now = datetime.now(CET)
            quotes = P.PROVIDERS[name](http, amounts)
            # Group like the sheet: one block per payment method, amounts ascending.
            order = {"bank transfer": 0, "debit card": 1, "credit card": 2}
            quotes.sort(key=lambda q: (order.get(q.method, 9), q.amount))
            kept = 0
            for q in quotes:
                q.mid_rate = q.mid_rate or mid
                if abs(q.rate - q.mid_rate) / q.mid_rate > MAX_DEVIATION:
                    print(f"  ! {q.provider} {q.method} €{q.amount:g}: rate {q.rate} is far from mid "
                          f"{q.mid_rate}; dropped", file=sys.stderr)
                    continue
                q.notes.append(f"Auto-scraped; mid rate = {mid_source}")
                rows.append((now, q))
                kept += 1
            if not kept:
                raise LookupError("no usable quotes returned")
            print(f"✓ {name}: {kept} quotes")
        except Exception as e:  # keep going: one broken provider must not lose the others
            failures.append((name, f"{type(e).__name__}: {e}"))
            print(f"✗ {name}: {type(e).__name__}: {e}", file=sys.stderr)
    return rows, failures


def last_data_row(ws) -> int:
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, 1).value not in (None, ""):
            return r
    return 1


def write_xlsx(path: Path, rows):
    wb = openpyxl.load_workbook(path)
    ws = wb["Rates"]
    template = last_data_row(ws)
    r = template
    for now, q in rows:
        r += 1
        values = {
            1: q.provider,
            2: now.strftime("%Y-%m-%d"),
            3: q.method,
            4: q.amount,
            5: round(q.mid_rate, 5),
            7: now.strftime("%H:%M"),
            8: q.recipient_gets,
            9: round(q.fee, 2),
            10: round(q.rate, 4),
            11: round(q.promo_rate, 4) if q.promo_rate else None,
            12: f'=IF(OR(E{r}="",J{r}=""),"",(E{r}-J{r})/E{r})',
            15: " | ".join(q.notes),
        }
        for col in range(1, 16):
            src, dst = ws.cell(template, col), ws.cell(r, col)
            if src.has_style:
                dst._style = copy.copy(src._style)
            if col in values:
                dst.value = values[col]
    # Extend the payment-method dropdown over the new rows.
    for dv in ws.data_validations.dataValidation:
        if "bank transfer" in str(dv.formula1):
            dv.sqref = openpyxl.worksheet.cell_range.MultiCellRange(f"C2:C{max(r, 131)}")
    # openpyxl drops cached formula results; make Excel/Sheets recompute the FX Margin column on open.
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)


def write_csv(path: Path, rows):
    new = not path.exists()
    with path.open("a", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=CSV_FIELDS)
        if new:
            w.writeheader()
        for now, q in rows:
            w.writerow({
                "provider": q.provider, "capture_date": now.strftime("%Y-%m-%d"),
                "capture_time_cet": now.strftime("%H:%M"), "payment_method": q.method,
                "amount_eur": q.amount, "mid_rate": round(q.mid_rate, 5), "recipient_inr": q.recipient_gets,
                "fee_eur": round(q.fee, 2), "rate": round(q.rate, 4),
                "promo_rate": round(q.promo_rate, 4) if q.promo_rate else "",
                "fx_margin": round((q.mid_rate - q.rate) / q.mid_rate, 6), "notes": " | ".join(q.notes),
            })


def summary(rows, failures) -> str:
    lines = ["| Provider | Method | € sent | ₹ received | Fee € | Rate | FX margin |",
             "|---|---|---:|---:|---:|---:|---:|"]
    for _, q in rows:
        lines.append(f"| {q.provider} | {q.method} | {q.amount:g} | {q.recipient_gets:,.2f} | {q.fee:.2f} "
                     f"| {q.rate:.4f} | {(q.mid_rate - q.rate) / q.mid_rate:.2%} |")
    if failures:
        lines += ["", "**Failed providers**", ""] + [f"- `{n}`: {err}" for n, err in failures]
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--providers", default="all", help=f"comma list from: {','.join(P.PROVIDERS)}")
    ap.add_argument("--amounts", default=",".join(map(str, AMOUNTS)))
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--csv", type=Path, default=DEFAULT_CSV)
    ap.add_argument("--dry-run", action="store_true", help="collect and print, write nothing")
    ap.add_argument("--debug", action="store_true", help="print raw API responses")
    args = ap.parse_args()

    P.DEBUG = args.debug
    names = list(P.PROVIDERS) if args.providers == "all" else [n.strip() for n in args.providers.split(",")]
    unknown = [n for n in names if n not in P.PROVIDERS]
    if unknown:
        ap.error(f"unknown providers: {unknown}")
    amounts = [int(a) for a in args.amounts.split(",")]

    rows, failures = collect(names, amounts)
    report = summary(rows, failures)
    print("\n" + report)
    if os.environ.get("GITHUB_STEP_SUMMARY"):
        with open(os.environ["GITHUB_STEP_SUMMARY"], "a", encoding="utf-8") as f:
            f.write("## EUR → INR scrape\n\n" + report + "\n")

    if rows and not args.dry_run:
        write_xlsx(args.xlsx, rows)
        write_csv(args.csv, rows)
        print(f"\nAppended {len(rows)} rows to {args.xlsx.name} and {args.csv.name}")
    sys.exit(1 if failures else 0)


if __name__ == "__main__":
    main()
